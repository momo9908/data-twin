"""材料表格读取、工程量换算和独立轮盘参数存储。

只读输入文件；转换全部有效点，保持试验行序，不做平滑、拟合或分段截断。
列号采用 0 基，文件数据起始行采用 1 基。
"""
from dataclasses import asdict, dataclass, field
import codecs
import csv
from collections import Counter
import io
import json
import math
import os
from pathlib import Path
import re
import tempfile

import numpy as np


class MaterialImportCancelled(Exception):
    """读取任务已取消。"""


@dataclass
class DiscParameters:
    outer_diameter_mm: float = 1.0
    circumferential_coefficient: float = 1.0
    radial_coefficient: float = 1.0


@dataclass
class MaterialPreview:
    path: str
    sheet_names: list
    sheet: object
    rows: list
    column_labels: list
    suggested_strain_column: int = 0
    suggested_stress_column: int = 1
    suggested_start_row: int = 2
    strain_unit: str = '无量纲'
    stress_unit: str = 'MPa'
    encoding: str = 'auto'
    delimiter: str = 'auto'
    source_rows: list = field(default_factory=list)


@dataclass
class MaterialCurve:
    path: str
    sheet: object
    engineering_strain: np.ndarray
    engineering_stress: np.ndarray
    true_strain: np.ndarray
    true_stress: np.ndarray
    source_rows: np.ndarray
    break_before: np.ndarray
    rejected_rows: list = field(default_factory=list)
    blank_rows: int = 0
    warnings: list = field(default_factory=list)

    @property
    def point_count(self):
        return len(self.engineering_strain)


@dataclass
class AutoMaterialImport:
    preview: MaterialPreview
    curve: MaterialCurve
    notes: list = field(default_factory=list)


@dataclass
class _Cell:
    value: object
    percent: bool = False


def _check_cancelled(cancelled):
    if cancelled is not None and cancelled():
        raise MaterialImportCancelled('已取消材料文件读取。')


def _validate_extension(path):
    suffix = Path(path).suffix.lower()
    if suffix == '.xls':
        raise ValueError('暂不支持旧版 .xls，请先另存为 .xlsx 再导入。')
    if suffix not in ('.xlsx', '.xlsm', '.csv', '.txt', '.tsv', '.dat'):
        raise ValueError('请选择 .xlsx、.xlsm、.csv、.txt、.tsv 或 .dat 表格文本。')
    return suffix


def _text_options(path, encoding='auto', delimiter='auto'):
    with open(path, 'rb') as stream:
        sample = stream.read(65536)
    if encoding == 'auto':
        if sample.startswith((codecs.BOM_UTF32_LE, codecs.BOM_UTF32_BE)):
            encoding = 'utf-32'
        elif sample.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
            encoding = 'utf-16'
        elif sample and sample.count(b'\x00') > len(sample) // 5:
            # BOM-less UTF-16 commonly occurs in instrument exports.
            even = sample[::2].count(b'\x00')
            odd = sample[1::2].count(b'\x00')
            encoding = 'utf-16-le' if odd > even else 'utf-16-be'
        else:
            for candidate in ('utf-8-sig', 'gb18030'):
                try:
                    codecs.getincrementaldecoder(candidate)().decode(sample, final=False)
                    encoding = candidate
                    break
                except UnicodeDecodeError:
                    continue
            else:
                raise ValueError('无法自动识别文本编码，请将文件保存为 UTF-8 或 Unicode 文本后导入。')
    try:
        decoded = codecs.getincrementaldecoder(encoding)().decode(sample, final=False)
    except (UnicodeError, LookupError) as exc:
        raise ValueError('所选文本编码无法读取文件，请重新选择编码。') from exc
    aliases = {'逗号': ',', '制表符': '\t', '分号': ';', '空白': 'whitespace'}
    delimiter = aliases.get(delimiter, delimiter)
    if delimiter == 'auto':
        # Score actual tabular rows rather than punctuation in title/remark lines.
        scores = []
        for separator in (',', '\t', ';', '|', 'whitespace'):
            try:
                rows = ([_split_whitespace_line(line) for line in decoded.splitlines()[:2000]]
                        if separator == 'whitespace' else
                        list(csv.reader(io.StringIO(decoded), delimiter=separator))[:2000])
            except csv.Error:
                continue
            numeric_widths = []
            header_score = 0
            for row in rows:
                count = 0
                for value in row:
                    try:
                        _number(_Cell(value))
                        count += 1
                    except (ValueError, TypeError, OverflowError):
                        pass
                if count >= 2:
                    numeric_widths.append(len(row))
                strain = [i for i, value in enumerate(row) if _material_role(value) == 'strain']
                stress = [i for i, value in enumerate(row) if _material_role(value) == 'stress']
                if strain and stress and any(a != b for a in strain for b in stress):
                    header_score = 10000
            common = Counter(numeric_widths).most_common(1)
            numeric_score = common[0][1] if common else 0
            explicit = separator != 'whitespace'
            scores.append((header_score + min(numeric_score, 2000), explicit, separator))
        delimiter = max(scores, key=lambda item: (item[0], item[1]))[2] if scores else 'whitespace'
    if delimiter not in (',', '\t', ';', '|', 'whitespace'):
        raise ValueError('不支持的分隔符。')
    return encoding, delimiter


def _display_cell(cell):
    value = cell.value
    if value is None:
        return ''
    if cell.percent and isinstance(value, (int, float)) and not isinstance(value, bool):
        return format(value * 100, '.12g') + '%'
    return str(value)


def _split_whitespace_line(line):
    text = line.strip()
    if re.search(r'\t| {2,}', text):
        return re.split(r'\t+| {2,}', text)
    return text.split()


def _rows(path, sheet=None, encoding='auto', delimiter='auto', cancelled=None):
    """产生 (源行号, Cell列表)，调用方必须关闭生成器以释放工作簿。"""
    suffix = _validate_extension(path)
    _check_cancelled(cancelled)
    if suffix in ('.xlsx', '.xlsm'):
        try:
            import openpyxl
        except ImportError as exc:
            raise ValueError('读取 Excel 需要当前环境中的 openpyxl。') from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            selected = workbook[sheet] if sheet is not None else workbook.active
            for row_number, row in enumerate(selected.iter_rows(), 1):
                _check_cancelled(cancelled)
                yield row_number, [_Cell(cell.value, '%' in (cell.number_format or '')) for cell in row]
        finally:
            workbook.close()
    else:
        encoding, delimiter = _text_options(path, encoding, delimiter)
        try:
            with open(path, 'r', encoding=encoding, newline='') as stream:
                if delimiter == 'whitespace':
                    for row_number, line in enumerate(stream, 1):
                        _check_cancelled(cancelled)
                        yield row_number, [_Cell(value) for value in _split_whitespace_line(line)]
                else:
                    reader = csv.reader(stream, delimiter=delimiter)
                    for row in reader:
                        _check_cancelled(cancelled)
                        yield reader.line_num, [_Cell(value) for value in row]
        except UnicodeError as exc:
            raise ValueError('文本编码不一致，无法完整读取；请统一保存为 UTF-8 后导入。') from exc


def _number(cell):
    value = cell.value
    if value is None or isinstance(value, bool):
        raise ValueError('缺少数值或单元格不是数值')
    percent = cell.percent
    if isinstance(value, str):
        value = value.strip().replace('\u2212', '-').replace('\uff0d', '-')
        if value.endswith(('%', '％')):
            value = value[:-1].strip()
            percent = True
            numeric = float(value) / 100.0
        else:
            numeric = float(value)
    else:
        numeric = float(value)
    if not math.isfinite(numeric):
        raise ValueError('数值不是有限数')
    return numeric, percent


def _material_role(value):
    """只识别工程/名义/未限定的应力应变，排除真实量及塑性应变。"""
    text = str(value or '').strip().lower()
    if any(token in text for token in ('true', '真实', '真应', 'plastic', '塑性')):
        return None
    strain = any(token in text for token in ('应变', 'strain', '延伸率', '伸长率', 'elongation'))
    stress = any(token in text for token in ('应力', 'stress'))
    compact = re.sub(r'[\s_{}()\[\]]', '', text)
    strain = strain or compact in ('e', 'ε', 'εe', 'epsilon', 'eps', 'en')
    stress = stress or compact in ('s', 'σ', 'σe', 'sigma', 'sn')
    if strain == stress:
        return None
    return 'strain' if strain else 'stress'


def _header_text(rows, index, column):
    """加入上一层 Engineering/True 分组标题，兼容 Excel 合并表头。"""
    text = str(rows[index][1][column].value or '')
    for _, previous in rows[max(0, index-2):index]:
        values = [str(cell.value or '') for cell in previous[:column+1]]
        context = next((value for value in reversed(values) if value.strip()), '')
        lowered = context.lower()
        if any(token in lowered for token in ('engineering', 'nominal', 'true', 'plastic',
                                               '工程', '名义', '真实', '塑性')):
            # Do not carry a previous column's fully named header into this column.
            if _material_role(context) is None and not any(
                    token in lowered for token in ('strain', 'stress', '应力', '应变')):
                text = context+' '+text
    return text


def _ranked_columns(rows, index, role):
    matches = []
    for column in range(len(rows[index][1])):
        text = _header_text(rows, index, column)
        if _material_role(text) == role:
            rank = 2 if any(token in text.lower() for token in
                            ('engineering', 'nominal', '工程', '名义')) else 1
            matches.append((rank, column))
    if not matches:
        return []
    best = max(rank for rank, _ in matches)
    return [column for rank, column in matches if rank == best]


def _numeric_pair(row, strain_column, stress_column):
    if max(strain_column, stress_column) >= len(row):
        return False
    try:
        _number(row[strain_column])
        _, stress_percent = _number(row[stress_column])
        return not stress_percent
    except (ValueError, TypeError, OverflowError):
        return False


def _automatic_unit(rows, column, kind, data_index, header_index):
    """单位只从标注或百分比格式读取；缺失时采用明确记录的约定。"""
    found = set()
    start = header_index if header_index >= 0 else max(0, data_index-3)
    for _, row in rows[start:data_index]:
        if column >= len(row):
            continue
        text = str(row[column].value or '').strip().lower()
        compact = re.sub(r'[\s()\[\]{}：:]', '', text).removeprefix('单位').removeprefix('unit')
        unit_only = compact in ('%', '％', 'percent', '百分数', '百分比', 'με', 'µε', '微应变',
                                'microstrain', '无量纲', 'dimensionless', 'mm/mm', 'm/m',
                                'mm', 'mpa', 'gpa', 'kpa', 'pa', 'n/mm2', 'n/mm²',
                                'n/mm^2', 'psi', 'ksi', 'kgf/mm2', 'kgf/mm²', 'bar')
        if _material_role(text) != kind and not unit_only:
            continue
        if kind == 'strain':
            if any(token in text for token in ('με', 'µε', 'microstrain', '微应变')):
                found.add('με')
            elif any(token in text for token in ('%', '％', '百分', 'percent')):
                found.add('%')
            elif any(token in text for token in ('无量纲', 'dimensionless', 'mm/mm', 'm/m')):
                found.add('无量纲')
            elif re.search(r'(^|[\s(\[])mm([\s)\]]|$)', text):
                raise ValueError('识别到应变列使用长度 mm；需要工程应变或延伸率，不能直接使用伸长量。')
        else:
            for token, unit in (('gpa', 'GPa'), ('kpa', 'kPa'), ('mpa', 'MPa')):
                if token in text:
                    found.add(unit)
                    break
            else:
                if re.search(r'(^|[^a-z])pa([^a-z]|$)', text):
                    found.add('Pa')
                elif re.search(r'n\s*/\s*mm\s*(?:\^?2|²)', text):
                    found.add('MPa')
                elif any(token in text for token in ('psi', 'ksi', 'kgf', 'bar')):
                    raise ValueError('应力列单位不是支持的 MPa、Pa、kPa 或 GPa，请统一单位后导入。')
    if kind == 'strain' and not found:
        for _, row in rows[data_index:data_index+20]:
            if column < len(row):
                try:
                    _, percent = _number(row[column])
                    if percent:
                        found.add('%')
                except (ValueError, TypeError, OverflowError):
                    pass
    if len(found) > 1:
        raise ValueError('同一%s列存在相互冲突的单位标注。' % ('应变' if kind == 'strain' else '应力'))
    if found:
        return next(iter(found)), None
    unit = '无量纲' if kind == 'strain' else 'MPa'
    return unit, ('应变单位未标注，按无量纲小数读取。' if kind == 'strain' else
                  '应力单位未标注，按 MPa 读取。')


def _detect_material_preview(path, sheet, sheet_names, encoding, delimiter, cancelled):
    rows = []
    iterator = _rows(path, sheet, encoding, delimiter, cancelled)
    try:
        for row_number, row in iterator:
            rows.append((row_number, row))
            if len(rows) >= 2000:
                break
    finally:
        iterator.close()
    if not rows:
        raise ValueError('没有可读取的数据行。')
    candidates = {}
    semantic_headers_found = False
    for index, (_, row) in enumerate(rows):
        _check_cancelled(cancelled)
        strain = _ranked_columns(rows, index, 'strain')
        stress = _ranked_columns(rows, index, 'stress')
        if not strain or not stress:
            continue
        semantic_headers_found = True
        pairs = [(e, s) for e in strain for s in stress if e != s and
                 sum(_numeric_pair(values, e, s) for _, values in rows[index+1:]) >= 2]
        if len(pairs) > 1:
            raise ValueError('检测到多组工程应力—应变列，无法唯一判断；请将所需材料单独保存成一个表格。')
        if pairs:
            candidates.setdefault(pairs[0], index)
    notes = []
    confidence = 2
    if len(candidates) > 1:
        raise ValueError('检测到不同的应力—应变数据分组，无法唯一判断；请每个文件保留一组材料数据。')
    if candidates:
        (strain_column, stress_column), header_index = next(iter(candidates.items()))
        possible = [(index, values) for index, (_, values) in enumerate(rows)
                    if index > header_index and _numeric_pair(values, strain_column, stress_column)]
    else:
        if semantic_headers_found:
            raise ValueError('检测到应力/应变表头，但标题列与数值列无法对应；请使用明确的逗号、制表符或双空格分列。')
        for _, row in rows:
            text = ' '.join(str(cell.value or '').lower() for cell in row)
            if (any(token in text for token in ('strain', '应变')) and
                    any(token in text for token in ('stress', '应力')) and
                    any(token in text for token in ('true', '真实', 'plastic', '塑性'))):
                raise ValueError('只识别到真实量或塑性量表头，未找到明确的工程应力与工程总应变。')
        numeric_rows = []
        for index, (_, row) in enumerate(rows):
            nonempty = [i for i, cell in enumerate(row)
                        if cell.value is not None and str(cell.value).strip()]
            if len(nonempty) == 2 and _numeric_pair(row, nonempty[0], nonempty[1]):
                numeric_rows.append((index, nonempty))
        pairs = Counter(tuple(columns) for _, columns in numeric_rows)
        if not pairs:
            raise ValueError('未找到工程应变/工程应力表头，且数据并非可确定的两列数值；请在文件中标明列名称。')
        if len(pairs) > 1:
            raise ValueError('文件中存在不同位置的两列数值，无法唯一判断材料数据。')
        (strain_column, stress_column), count = pairs.most_common(1)[0]
        if count < 2:
            raise ValueError('至少需要两行有效的工程应力—应变数据。')
        header_index = -1
        confidence = 1
        notes.append('无明确表头：按两列数据约定，前列为工程应变，后列为工程应力。')
        matching_indices = {index for index, columns in numeric_rows
                            if tuple(columns) == (strain_column, stress_column)}
        possible = [(index, row) for index, (_, row) in enumerate(rows) if index in matching_indices]
    data_index = possible[0][0]
    strain_unit, strain_note = _automatic_unit(rows, strain_column, 'strain', data_index, header_index)
    stress_unit, stress_note = _automatic_unit(rows, stress_column, 'stress', data_index, header_index)
    notes.extend(note for note in (strain_note, stress_note) if note)
    start = max(0, header_index if header_index >= 0 else data_index-1)
    displayed = rows[start:start+20]
    width = max(len(row) for _, row in displayed)
    header = rows[header_index][1] if header_index >= 0 else []
    labels = ['第 %d 列%s' % (i+1, '：'+_display_cell(header[i])[:60]
                             if i < len(header) and _display_cell(header[i]) else '') for i in range(width)]
    preview = MaterialPreview(str(Path(path).resolve()), list(sheet_names), sheet,
                              [[_display_cell(cell) for cell in row] for _, row in displayed],
                              labels, strain_column, stress_column, rows[data_index][0],
                              strain_unit, stress_unit, encoding, delimiter,
                              [row_number for row_number, _ in displayed])
    return preview, notes, confidence


def auto_import_material(path, cancelled=None):
    """自动识别一组材料数据并全量换算；信息不足或多组歧义时明确停止。"""
    path = str(Path(path).resolve())
    suffix = _validate_extension(path)
    _check_cancelled(cancelled)
    sheet_names = []
    if suffix in ('.xlsx', '.xlsm'):
        try:
            import openpyxl
        except ImportError as exc:
            raise ValueError('读取 Excel 需要当前环境中的 openpyxl。') from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            sheet_names = list(workbook.sheetnames)
        finally:
            workbook.close()
        encoding, delimiter = 'auto', 'auto'
        sheets = sheet_names
    else:
        encoding, delimiter = _text_options(path)
        sheets = [None]
    candidates, errors = [], []
    for sheet in sheets:
        _check_cancelled(cancelled)
        try:
            candidates.append(_detect_material_preview(path, sheet, sheet_names, encoding, delimiter, cancelled))
        except ValueError as exc:
            errors.append((sheet, str(exc)))
    if not candidates:
        detail = '；'.join((str(sheet)+'：' if sheet else '')+error for sheet, error in errors[:4])
        raise ValueError('无法自动识别工程应力—应变数据。'+detail)
    confidence = max(candidate[2] for candidate in candidates)
    candidates = [candidate for candidate in candidates if candidate[2] == confidence]
    if len(candidates) != 1:
        raise ValueError('多个工作表包含同等明确的材料数据（%s），无法唯一判断；请将需要的材料单独保存后导入。' %
                         '、'.join(str(candidate[0].sheet) for candidate in candidates))
    preview, notes, _ = candidates[0]
    _check_cancelled(cancelled)
    curve = load_material_curve(path, preview.suggested_strain_column, preview.suggested_stress_column,
                                preview.suggested_start_row, preview.strain_unit, preview.stress_unit,
                                preview.sheet, preview.encoding, preview.delimiter, cancelled)
    return AutoMaterialImport(preview, curve, notes)


def inspect_material_file(path, sheet=None, encoding='auto', delimiter='auto',
                          preview_rows=20, cancelled=None):
    path = str(Path(path).resolve())
    suffix = _validate_extension(path)
    sheet_names = []
    if suffix in ('.xlsx', '.xlsm'):
        try:
            import openpyxl
        except ImportError as exc:
            raise ValueError('读取 Excel 需要当前环境中的 openpyxl。') from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            sheet_names = list(workbook.sheetnames)
            sheet = sheet or workbook.active.title
            if sheet not in sheet_names:
                raise ValueError('所选工作表不存在。')
        finally:
            workbook.close()
    else:
        encoding, delimiter = _text_options(path, encoding, delimiter)
    captured = []
    iterator = _rows(path, sheet, encoding, delimiter, cancelled)
    try:
        for _, row in iterator:
            captured.append(row)
            if len(captured) >= max(1, int(preview_rows)):
                break
    finally:
        iterator.close()
    if not captured or not any(any(cell.value is not None and str(cell.value).strip()
                                   for cell in row) for row in captured):
        raise ValueError('文件或所选工作表没有可预览的数据。')
    width = max(len(row) for row in captured)
    labels = ['第 %d 列' % (i + 1) for i in range(width)]
    strain_col, stress_col, header_row = 0, min(1, width - 1), -1
    for row_index, row in enumerate(captured):
        strain_matches = [i for i, cell in enumerate(row)
                          if any(key in str(cell.value).lower() for key in ('应变', 'strain', 'elongation'))]
        stress_matches = [i for i, cell in enumerate(row)
                          if any(key in str(cell.value).lower() for key in ('应力', 'stress'))]
        if strain_matches and stress_matches:
            strain_col, stress_col, header_row = strain_matches[0], stress_matches[0], row_index
            labels = ['第 %d 列：%s' % (i + 1, _display_cell(row[i])[:60])
                      if i < len(row) and _display_cell(row[i]) else '第 %d 列' % (i + 1)
                      for i in range(width)]
            break
    strain_unit, stress_unit = '无量纲', 'MPa'
    for row in captured[:max(2, header_row + 2)]:
        if strain_col < len(row):
            text = str(row[strain_col].value).lower()
            if any(unit in text for unit in ('με', 'µε', 'microstrain', '微应变')):
                strain_unit = 'με'
            elif any(unit in text for unit in ('%', '％', '百分')):
                strain_unit = '%'
        if stress_col < len(row):
            text = str(row[stress_col].value).lower()
            for token, unit in (('gpa', 'GPa'), ('kpa', 'kPa'), ('mpa', 'MPa')):
                if token in text:
                    stress_unit = unit
                    break
            else:
                if re.search(r'(^|[^a-z])pa([^a-z]|$)', text):
                    stress_unit = 'Pa'
    start_row = max(1, header_row + 2)
    for index, row in enumerate(captured):
        if index <= header_row or max(strain_col, stress_col) >= len(row):
            continue
        try:
            _number(row[strain_col])
            _number(row[stress_col])
            start_row = index + 1
            break
        except (ValueError, TypeError, OverflowError):
            continue
    return MaterialPreview(path, sheet_names, sheet,
                           [[_display_cell(cell) for cell in row] for row in captured],
                           labels, strain_col, stress_col, start_row,
                           strain_unit, stress_unit, encoding, delimiter)


def load_material_curve(path, strain_column, stress_column, start_row=2,
                        strain_unit='无量纲', stress_unit='MPa', sheet=None,
                        encoding='auto', delimiter='auto', cancelled=None):
    if strain_column < 0 or stress_column < 0 or strain_column == stress_column:
        raise ValueError('请为工程应变和工程应力选择不同的数据列。')
    if int(start_row) < 1:
        raise ValueError('数据起始行必须至少为 1。')
    strain_scales = {'无量纲': 1.0, '%': 0.01, 'με': 1e-6}
    stress_scales = {'MPa': 1.0, 'Pa': 1e-6, 'kPa': 0.001, 'GPa': 1000.0}
    if strain_unit not in strain_scales or stress_unit not in stress_scales:
        raise ValueError('请选择支持的应变与应力单位。')
    path = str(Path(path).resolve())
    eng_e, eng_s, true_e, true_s, source_rows, breaks = [], [], [], [], [], []
    rejected, blank_rows, pending_break = [], 0, False
    iterator = _rows(path, sheet, encoding, delimiter, cancelled)
    try:
        for row_number, row in iterator:
            if row_number < int(start_row):
                continue
            if not row or all(cell.value is None or str(cell.value).strip() == '' for cell in row):
                blank_rows += 1
                pending_break = bool(eng_e)
                continue
            try:
                if max(strain_column, stress_column) >= len(row):
                    raise ValueError('该行缺少选定的数据列')
                e, explicit_percent = _number(row[strain_column])
                s, stress_percent = _number(row[stress_column])
                if stress_percent:
                    raise ValueError('应力数据不能使用百分数单位')
                e *= 1.0 if explicit_percent else strain_scales[strain_unit]
                s *= stress_scales[stress_unit]
                if not math.isfinite(e) or not math.isfinite(s):
                    raise ValueError('单位换算后不是有限数')
                if e <= -1.0:
                    raise ValueError('工程应变必须大于 -1，才能计算 ln(1+工程应变)')
                with np.errstate(over='ignore', invalid='ignore', divide='ignore'):
                    t_e = float(np.log1p(e))
                    t_s = float(np.float64(s) * (np.float64(e) + 1.0))
                if not math.isfinite(t_e) or not math.isfinite(t_s):
                    raise ValueError('公式计算结果不是有限数')
            except (ValueError, TypeError, OverflowError) as exc:
                rejected.append((row_number, str(exc)))
                pending_break = bool(eng_e)
                continue
            breaks.append(pending_break)
            pending_break = False
            eng_e.append(e)
            eng_s.append(s)
            true_e.append(t_e)
            true_s.append(t_s)
            source_rows.append(row_number)
    finally:
        iterator.close()
    _check_cancelled(cancelled)
    if len(eng_e) < 2:
        detail = '；首个问题：第 %d 行，%s' % rejected[0] if rejected else ''
        raise ValueError('至少需要两行有效的工程应力—应变数据'+detail)
    warnings = []
    if rejected:
        warnings.append('跳过 %d 行无效数据，曲线在缺失位置断开。' % len(rejected))
    if blank_rows:
        warnings.append('跳过 %d 个空行。' % blank_rows)
    return MaterialCurve(path, sheet, np.asarray(eng_e), np.asarray(eng_s),
                         np.asarray(true_e), np.asarray(true_s), np.asarray(source_rows, dtype=int),
                         np.asarray(breaks, dtype=bool), rejected, blank_rows, warnings)


def validate_disc_parameters(parameters):
    values = asdict(parameters)
    for key, value in values.items():
        if isinstance(value, bool):
            raise ValueError('轮盘参数必须为有限数值。')
        try:
            values[key] = float(value)
        except (ValueError, TypeError, OverflowError) as exc:
            raise ValueError('轮盘参数必须为有限数值。') from exc
        if not math.isfinite(values[key]):
            raise ValueError('轮盘参数必须为有限数值。')
    if values['outer_diameter_mm'] <= 0:
        raise ValueError('轮盘外径必须大于 0 mm。')
    return DiscParameters(**values)


def load_disc_parameters(path):
    path = Path(path)
    if not path.exists():
        return DiscParameters()
    try:
        payload = json.loads(path.read_text(encoding='utf-8-sig'))
        if not isinstance(payload, dict):
            raise ValueError('配置应为对象')
        values = {key: payload.get(key, default) for key, default in asdict(DiscParameters()).items()}
        return validate_disc_parameters(DiscParameters(**values))
    except (ValueError, TypeError, OverflowError) as exc:
        raise ValueError('轮盘参数配置无效：'+str(exc)) from exc


def save_disc_parameters(path, parameters):
    parameters = validate_disc_parameters(parameters)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(mode='w', encoding='utf-8', suffix='.tmp',
                                         prefix=path.name+'.', dir=path.parent, delete=False) as stream:
            temporary = stream.name
            json.dump(asdict(parameters), stream, ensure_ascii=False, indent=2, allow_nan=False)
            stream.write('\n')
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        if temporary is not None and os.path.exists(temporary):
            os.unlink(temporary)


def export_material_csv(path, curve):
    destination = Path(path).resolve()
    source = Path(curve.path).resolve()
    same_file = os.path.normcase(str(destination)) == os.path.normcase(str(source))
    if destination.exists() and source.exists():
        same_file = same_file or os.path.samefile(destination, source)
    if same_file:
        raise ValueError('导出文件不能覆盖导入的原始材料文件，请选择其他文件名。')
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(mode='w', encoding='utf-8-sig', newline='',
                                         suffix='.tmp', prefix=destination.name+'.',
                                         dir=destination.parent, delete=False) as stream:
            temporary = stream.name
            writer = csv.writer(stream)
            writer.writerow(['工程应变', '工程应力（MPa）', '真实应变', '真实应力（MPa）'])
            for row in zip(curve.engineering_strain, curve.engineering_stress,
                           curve.true_strain, curve.true_stress):
                writer.writerow([format(float(value), '.17g') for value in row])
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, destination)
    finally:
        if temporary is not None and os.path.exists(temporary):
            os.unlink(temporary)
    return str(destination)
